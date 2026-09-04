from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from conftest import MemorySecretStore
from docx import Document as DocxDocument
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, RectangleObject

from careerflow_agent.config import Settings
from careerflow_agent.contracts import CandidateProfileInput
from careerflow_agent.database import Database
from careerflow_agent.document_parser import extract_profile_suggestions, parse_resume
from careerflow_agent.profile_vault import (
    PROFILE_KEY_REFERENCE,
    ProfileKeyUnavailableError,
    ProfileVault,
    ProfileVersionConflictError,
    ResumeAlreadyImportedError,
)

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def resume_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Experience", level=1)
    document.add_paragraph("Built a deterministic workflow with typed contracts.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def image_only_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def selectable_text_pdf() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = writer._add_object(  # noqa: SLF001 - pypdf's low-level writer builds a tiny fixture.
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Built a local supervised agent.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)  # noqa: SLF001
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def linked_profile_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_uri(
        page_number=0,
        uri="https://www.linkedin.com/in/synthetic-candidate?tracking=fixture",
        rect=RectangleObject((72, 700, 160, 720)),
    )
    writer.add_uri(
        page_number=0,
        uri="http://github.com/synthetic-candidate#profile",
        rect=RectangleObject((170, 700, 250, 720)),
    )
    writer.add_uri(
        page_number=0,
        uri="https://github.com/synthetic-candidate/example-project",
        rect=RectangleObject((260, 700, 350, 720)),
    )
    writer.add_uri(
        page_number=0,
        uri="javascript:alert('fixture')",
        rect=RectangleObject((360, 700, 430, 720)),
    )
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def profile_input() -> CandidateProfileInput:
    return CandidateProfileInput(
        first_name="Synthetic",
        last_name="Candidate",
        email="synthetic.candidate@example.com",
        city="Test City",
        region="CA",
        country="United States",
        school="Example University",
        degree="Bachelor of Science",
        field_of_study="Computer Science",
        graduation_year=2027,
        work_authorization="authorized",
        requires_sponsorship=False,
    )


@pytest.mark.asyncio
async def test_profile_versions_are_encrypted_and_restored(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, local_token="test-token", telemetry_enabled=False)
    database = Database(settings.database_url)
    secrets = MemorySecretStore()
    vault = ProfileVault(database, secrets, tmp_path)
    await database.initialize()
    try:
        first = await vault.save(profile_input(), expected_version=None)
        second_input = profile_input().model_copy(update={"preferred_name": "Synth"})
        second = await vault.save(second_input, expected_version=1)
        restored = await vault.load()

        assert first.profile.version == 1
        assert second.profile.version == 2
        assert restored is not None
        assert restored.profile.id == first.profile.id
        assert restored.profile.version == 2
        assert any(
            fact.canonical_path == "identity.preferred_name" and fact.value == "Synth"
            for fact in restored.profile.facts
        )
        database_bytes = (tmp_path / "careerflow.db").read_bytes()
        assert b"synthetic.candidate@example.com" not in database_bytes
        assert b"Example University" not in database_bytes
        assert PROFILE_KEY_REFERENCE in secrets.values
    finally:
        await database.close()


@pytest.mark.asyncio
async def test_profile_save_rejects_stale_version(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, local_token="test-token", telemetry_enabled=False)
    database = Database(settings.database_url)
    vault = ProfileVault(database, MemorySecretStore(), tmp_path)
    await database.initialize()
    try:
        await vault.save(profile_input(), expected_version=None)
        with pytest.raises(ProfileVersionConflictError):
            await vault.save(profile_input(), expected_version=None)
    finally:
        await database.close()


@pytest.mark.asyncio
async def test_missing_key_fails_closed(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, local_token="test-token", telemetry_enabled=False)
    database = Database(settings.database_url)
    secrets = MemorySecretStore()
    vault = ProfileVault(database, secrets, tmp_path)
    await database.initialize()
    try:
        await vault.save(profile_input(), expected_version=None)
        secrets.delete(PROFILE_KEY_REFERENCE)
        with pytest.raises(ProfileKeyUnavailableError):
            await vault.load()
    finally:
        await database.close()


@pytest.mark.asyncio
async def test_resume_import_is_encrypted_provenanced_reviewable_and_preserved(
    tmp_path: Path,
) -> None:
    settings = Settings(data_dir=tmp_path, local_token="test-token", telemetry_enabled=False)
    database = Database(settings.database_url)
    vault = ProfileVault(database, MemorySecretStore(), tmp_path)
    await database.initialize()
    content = resume_docx()
    try:
        first = await vault.save(profile_input(), expected_version=None)
        imported = await vault.import_resume(
            content,
            filename="Synthetic Resume.docx",
            media_type=DOCX_MEDIA_TYPE,
            expected_version=first.profile.version,
        )

        assert imported.snapshot.profile.version == 2
        assert imported.document.status == "parsed"
        assert imported.document.extracted_evidence_count == 3
        imported_evidence = [
            item
            for item in imported.snapshot.profile.evidence
            if item.source_document_id == imported.document.id
        ]
        assert len(imported_evidence) == 3
        assert all(not item.verified for item in imported_evidence)
        assert all(item.source_span is not None for item in imported_evidence)
        assert "paragraph" in (imported_evidence[0].source_span.section or "")

        artifact_path = tmp_path / imported.document.encrypted_artifact_ref
        encrypted_bytes = artifact_path.read_bytes()
        assert b"deterministic workflow" not in encrypted_bytes
        assert b"FastAPI" not in encrypted_bytes
        assert b"deterministic workflow" not in (tmp_path / "careerflow.db").read_bytes()

        reviewed = await vault.set_evidence_verification(
            imported_evidence[1].id,
            verified=True,
            expected_version=imported.snapshot.profile.version,
        )
        assert reviewed.profile.version == 3
        assert next(
            item for item in reviewed.profile.evidence if item.id == imported_evidence[1].id
        ).verified

        edited_input = profile_input().model_copy(update={"preferred_name": "Synth"})
        edited = await vault.save(edited_input, expected_version=reviewed.profile.version)
        assert edited.profile.version == 4
        assert edited.profile.source_documents == reviewed.profile.source_documents
        assert any(item.id == imported_evidence[1].id for item in edited.profile.evidence)

        with pytest.raises(ResumeAlreadyImportedError):
            await vault.import_resume(
                content,
                filename="Duplicate.docx",
                media_type=DOCX_MEDIA_TYPE,
                expected_version=edited.profile.version,
            )
    finally:
        await database.close()


@pytest.mark.asyncio
async def test_image_only_pdf_is_encrypted_and_marked_for_ocr(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, local_token="test-token", telemetry_enabled=False)
    database = Database(settings.database_url)
    vault = ProfileVault(database, MemorySecretStore(), tmp_path)
    await database.initialize()
    try:
        first = await vault.save(profile_input(), expected_version=None)
        imported = await vault.import_resume(
            image_only_pdf(),
            filename="Scanned Resume.pdf",
            media_type="application/pdf",
            expected_version=first.profile.version,
        )
        assert imported.document.status == "needs_ocr"
        assert imported.document.parse_error_code == "selectable_text_unavailable"
        assert imported.document.extracted_evidence_count == 0
        assert (tmp_path / imported.document.encrypted_artifact_ref).exists()
    finally:
        await database.close()


@pytest.mark.asyncio
async def test_selectable_pdf_creates_page_provenance(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path, local_token="test-token", telemetry_enabled=False)
    database = Database(settings.database_url)
    vault = ProfileVault(database, MemorySecretStore(), tmp_path)
    await database.initialize()
    try:
        first = await vault.save(profile_input(), expected_version=None)
        imported = await vault.import_resume(
            selectable_text_pdf(),
            filename="Text Resume.pdf",
            media_type="application/pdf",
            expected_version=first.profile.version,
        )
        evidence = next(
            item
            for item in imported.snapshot.profile.evidence
            if item.source_document_id == imported.document.id
        )
        assert imported.document.status == "parsed"
        assert evidence.statement == "Built a local supervised agent."
        assert evidence.source_span is not None
        assert evidence.source_span.page == 1
        assert evidence.source_span.section == "page:1"
    finally:
        await database.close()


def test_pdf_hyperlink_annotations_extract_profile_links_with_provenance() -> None:
    statements = parse_resume(linked_profile_pdf(), "application/pdf")
    suggestions = {
        suggestion.canonical_path: suggestion
        for suggestion in extract_profile_suggestions(statements)
    }

    assert suggestions["links.linkedin"].value == (
        "https://www.linkedin.com/in/synthetic-candidate"
    )
    assert suggestions["links.github"].value == "https://github.com/synthetic-candidate"
    assert suggestions["links.linkedin"].source_span.page == 1
    assert suggestions["links.github"].source_span.page == 1
    assert suggestions["links.linkedin"].source_span.section == "page:1 · hyperlink"
    assert suggestions["links.github"].source_span.section == "page:1 · hyperlink"
    assert all("example-project" not in statement.statement for statement in statements)
    assert all("javascript:" not in statement.statement for statement in statements)
