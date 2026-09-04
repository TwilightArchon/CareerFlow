from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import httpx
import pytest
from docx import Document as DocxDocument

from careerflow_agent.contracts import FieldMapping, JobPosting, ObservedForm
from careerflow_agent.job_ingestion import JobIngestionService, parse_job_page

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def resume_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Projects", level=1)
    document.add_paragraph("Built an observable supervised agent.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def material_resume_docx() -> bytes:
    document = DocxDocument()
    document.add_heading("Projects", level=1)
    document.add_paragraph("Built an observable supervised agent with tracing.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def profile_resume_docx() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Synthetic Candidate")
    document.add_paragraph(
        "synthetic.candidate@example.com | (404) 555-0123 | "
        "linkedin.com/in/synthetic-candidate | github.com/synthetic-candidate"
    )
    document.add_heading("Education", level=1)
    document.add_paragraph("Example University")
    document.add_paragraph("Bachelor of Science in Computer Science | Expected May 2027")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def test_health_requires_bearer_token(client: httpx.AsyncClient) -> None:
    response = await client.get("/v1/health")
    assert response.status_code == 200
    assert response.json()["databaseReady"] is True

    response = await client.get("/v1/health", headers={"Authorization": "Bearer wrong"})
    assert response.status_code == 401


async def test_profile_is_saved_versioned_and_restored(client: httpx.AsyncClient) -> None:
    empty = await client.get("/v1/profile")
    assert empty.status_code == 200
    assert empty.json() is None

    request = {
        "expected_version": None,
        "profile": {
            "first_name": "Synthetic",
            "last_name": "Candidate",
            "email": "synthetic.candidate@example.com",
            "city": "Test City",
            "region": "CA",
            "country": "United States",
            "school": "Example University",
            "degree": "Bachelor of Science",
            "field_of_study": "Computer Science",
            "graduation_year": 2027,
            "work_authorization": "authorized",
            "requires_sponsorship": False,
        },
    }
    saved = await client.put("/v1/profile", json=request)
    assert saved.status_code == 200
    assert saved.json()["profile"]["version"] == 1
    assert all(fact["state"] == "verified" for fact in saved.json()["profile"]["facts"])

    restored = await client.get("/v1/profile")
    assert restored.status_code == 200
    assert restored.json()["profile"] == saved.json()["profile"]

    stale = await client.put("/v1/profile", json=request)
    assert stale.status_code == 409


async def test_resume_import_and_evidence_review_are_authenticated_and_versioned(
    client: httpx.AsyncClient,
) -> None:
    resume = resume_docx()
    created = await client.put(
        "/v1/profile",
        json={
            "expected_version": None,
            "profile": {
                "first_name": "Synthetic",
                "last_name": "Candidate",
                "email": "synthetic.candidate@example.com",
            },
        },
    )
    assert created.status_code == 200

    imported = await client.post(
        "/v1/profile/resume",
        data={"expected_version": "1"},
        files={"file": ("Resume.docx", resume, DOCX_MEDIA_TYPE)},
    )
    assert imported.status_code == 201
    payload = imported.json()
    assert payload["snapshot"]["profile"]["version"] == 2
    assert payload["document"]["status"] == "parsed"
    assert payload["document"]["extractedEvidenceCount"] == 2
    imported_evidence = [
        item
        for item in payload["snapshot"]["profile"]["evidence"]
        if item["sourceDocumentId"] == payload["document"]["id"]
    ]
    assert all(not item["verified"] for item in imported_evidence)

    reviewed = await client.put(
        f"/v1/profile/evidence/{imported_evidence[1]['id']}/verification",
        json={"expected_version": 2, "verified": True},
    )
    assert reviewed.status_code == 200
    assert reviewed.json()["profile"]["version"] == 3
    assert next(
        item
        for item in reviewed.json()["profile"]["evidence"]
        if item["id"] == imported_evidence[1]["id"]
    )["verified"]

    duplicate = await client.post(
        "/v1/profile/resume",
        data={"expected_version": "3"},
        files={"file": ("Resume.docx", resume, DOCX_MEDIA_TYPE)},
    )
    assert duplicate.status_code == 409


async def test_resume_preview_extracts_profile_suggestions_without_creating_profile(
    client: httpx.AsyncClient,
) -> None:
    preview = await client.post(
        "/v1/profile/resume/preview",
        files={"file": ("Resume.docx", profile_resume_docx(), DOCX_MEDIA_TYPE)},
    )

    assert preview.status_code == 200
    payload = preview.json()
    assert payload["status"] == "parsed"
    assert payload["extractedEvidenceCount"] == 5
    values = {
        suggestion["canonicalPath"]: suggestion["value"] for suggestion in payload["suggestions"]
    }
    assert values == {
        "identity.first_name": "Synthetic",
        "identity.last_name": "Candidate",
        "contact.email": "synthetic.candidate@example.com",
        "contact.phone": "(404) 555-0123",
        "links.linkedin": "https://linkedin.com/in/synthetic-candidate",
        "links.github": "https://github.com/synthetic-candidate",
        "education.0.school": "Example University",
        "education.0.degree": "Bachelor of Science",
        "education.0.field_of_study": "Computer Science",
        "education.0.graduation_year": "2027",
    }
    assert (await client.get("/v1/profile")).json() is None


async def test_job_ingestion_endpoint_persists_reviewable_result(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    async def ingest_fixture(self: JobIngestionService, raw_url: str) -> JobPosting:
        metadata = (
            "Location: Raleigh, NC Job ID: R123 Company Name: Example Energy "
            "Profession: Engineering Job Description: Qualifications: "
            "Ability to build reliable software."
        )
        posting = parse_job_page(
            raw_url,
            "https://tenant.myworkdayjobs.com/jobs/123",
            f"""<html><head>
            <meta property="og:title" content="Software Engineering Intern">
            <meta property="og:description" content="{metadata}">
            </head></html>""",
        )
        return await self.database.save_job_posting(posting)

    monkeypatch.setattr(JobIngestionService, "ingest", ingest_fixture)
    ingested = await client.post(
        "/v1/jobs/ingest",
        json={"url": "https://tenant.myworkdayjobs.com/jobs/123?utm_source=fixture"},
    )
    assert ingested.status_code == 200
    payload = ingested.json()
    assert payload["title"] == "Software Engineering Intern"
    assert payload["company"] == "Example Energy"
    assert payload["platform"]["platform"] == "workday"

    restored = await client.get(f"/v1/jobs/{payload['id']}")
    assert restored.status_code == 200
    assert restored.json() == payload


async def test_prepare_materials_maps_only_verified_evidence(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    await client.put(
        "/v1/profile",
        json={
            "expected_version": None,
            "profile": {
                "first_name": "Synthetic",
                "last_name": "Candidate",
                "email": "synthetic.candidate@example.com",
            },
        },
    )
    imported = await client.post(
        "/v1/profile/resume",
        data={"expected_version": "1"},
        files={"file": ("Resume.docx", material_resume_docx(), DOCX_MEDIA_TYPE)},
    )
    evidence = next(
        item
        for item in imported.json()["snapshot"]["profile"]["evidence"]
        if item["sourceDocumentId"] is not None and "observable" in item["statement"]
    )
    verified = await client.put(
        f"/v1/profile/evidence/{evidence['id']}/verification",
        json={"expected_version": 2, "verified": True},
    )
    profile = verified.json()["profile"]

    async def ingest_fixture(self: JobIngestionService, raw_url: str) -> JobPosting:
        posting = parse_job_page(
            raw_url,
            raw_url,
            """<html><head><meta property="og:title" content="Agent Engineer"></head>
            <body><h2>Requirements</h2><ul>
            <li>Build an observable supervised agent.</li>
            </ul></body></html>""",
        )
        return await self.database.save_job_posting(posting)

    monkeypatch.setattr(JobIngestionService, "ingest", ingest_fixture)
    job = (
        await client.post(
            "/v1/jobs/ingest", json={"url": "https://example.com/jobs/agent-engineer"}
        )
    ).json()
    prepared = await client.post(
        "/v1/materials/prepare",
        json={
            "job_id": job["id"],
            "candidate_profile_id": profile["id"],
            "candidate_profile_version": profile["version"],
        },
    )

    assert prepared.status_code == 200
    payload = prepared.json()
    assert payload["modelUsed"] is False
    assert payload["supportedCount"] == 1
    assert payload["resumeDraft"]["entries"][0]["evidenceId"] == evidence["id"]
    assert payload["resumeDraft"]["entries"][0]["statement"] == (
        "Built an observable supervised agent with tracing."
    )
    assert b"supervised agent with tracing" not in (tmp_path / "careerflow.db").read_bytes()

    stale = await client.post(
        "/v1/materials/prepare",
        json={
            "job_id": job["id"],
            "candidate_profile_id": profile["id"],
            "candidate_profile_version": profile["version"] - 1,
        },
    )
    assert stale.status_code == 409


async def test_run_transition_is_idempotent(client: httpx.AsyncClient) -> None:
    created = await client.post(
        "/v1/runs",
        json={
            "job_id": str(uuid4()),
            "candidate_profile_id": str(uuid4()),
            "candidate_profile_version": 1,
            "auto_submit_authorized": False,
            "job_url": "https://example.com/jobs/1",
            "start_immediately": False,
        },
    )
    assert created.status_code == 201
    run_id = created.json()["id"]
    request = {
        "to_state": "ingesting_job",
        "reason_code": "run_started",
        "idempotency_key": "run-started-key-0001",
    }
    first = await client.post(f"/v1/runs/{run_id}/transitions", json=request)
    second = await client.post(f"/v1/runs/{run_id}/transitions", json=request)
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["id"] == second.json()["id"]


async def test_runs_are_listed_by_recent_activity_with_current_state(
    client: httpx.AsyncClient,
) -> None:
    urls = ["https://example.com/jobs/old", "https://example.com/jobs/new"]
    created_runs: list[dict[str, object]] = []
    for url in urls:
        response = await client.post(
            "/v1/runs",
            json={
                "job_id": str(uuid4()),
                "candidate_profile_id": str(uuid4()),
                "candidate_profile_version": 1,
                "auto_submit_authorized": False,
                "job_url": url,
                "start_immediately": False,
            },
        )
        assert response.status_code == 201
        created_runs.append(response.json())

    transitioned = await client.post(
        f"/v1/runs/{created_runs[0]['id']}/transitions",
        json={
            "to_state": "ingesting_job",
            "reason_code": "test_restore",
            "idempotency_key": "restore-transition-key-0001",
        },
    )
    assert transitioned.status_code == 200

    listed = await client.get("/v1/runs")
    assert listed.status_code == 200
    payload = listed.json()
    assert [run["id"] for run in payload] == [
        created_runs[0]["id"],
        created_runs[1]["id"],
    ]
    assert payload[0]["state"] == "ingesting_job"
    assert payload[0]["jobUrl"] == urls[0]
    assert payload[0]["createdAt"].endswith("Z")
    assert payload[0]["updatedAt"].endswith("Z")


async def test_user_confirmed_outcomes_are_append_only_and_drive_statistics(
    client: httpx.AsyncClient,
) -> None:
    run_ids: list[str] = []
    for suffix in ("one", "two"):
        created = await client.post(
            "/v1/runs",
            json={
                "job_id": str(uuid4()),
                "candidate_profile_id": str(uuid4()),
                "candidate_profile_version": 1,
                "auto_submit_authorized": False,
                "job_url": f"https://example.com/jobs/{suffix}",
                "start_immediately": False,
            },
        )
        assert created.status_code == 201
        run_ids.append(created.json()["id"])

    missing_confirmation = await client.put(
        f"/v1/runs/{run_ids[0]}/outcome",
        json={
            "outcome": "submitted",
            "reason_code": "user_confirmed_submitted",
            "confirmed_by_user": False,
        },
    )
    assert missing_confirmation.status_code == 422

    mismatched_reason = await client.put(
        f"/v1/runs/{run_ids[0]}/outcome",
        json={
            "outcome": "submitted",
            "reason_code": "validation_failed",
            "confirmed_by_user": True,
        },
    )
    assert mismatched_reason.status_code == 422

    submitted = await client.put(
        f"/v1/runs/{run_ids[0]}/outcome",
        json={
            "outcome": "submitted",
            "reason_code": "user_confirmed_submitted",
            "confirmed_by_user": True,
        },
    )
    assert submitted.status_code == 200
    assert submitted.json()["revision"] == 1
    assert submitted.json()["confirmation"][0]["kind"] == "user_correction"
    assert len(submitted.json()["confirmation"][0]["fingerprint"]) == 64

    duplicate = await client.put(
        f"/v1/runs/{run_ids[0]}/outcome",
        json={
            "outcome": "submitted",
            "reason_code": "user_confirmed_submitted",
            "confirmed_by_user": True,
        },
    )
    assert duplicate.json()["id"] == submitted.json()["id"]

    corrected = await client.put(
        f"/v1/runs/{run_ids[0]}/outcome",
        json={
            "outcome": "failed",
            "reason_code": "validation_failed",
            "confirmed_by_user": True,
        },
    )
    assert corrected.status_code == 200
    assert corrected.json()["revision"] == 2
    assert corrected.json()["supersedesId"] == submitted.json()["id"]
    assert corrected.json()["confirmation"] == []

    history = await client.get(f"/v1/runs/{run_ids[0]}/outcomes")
    assert history.status_code == 200
    assert [item["outcome"] for item in history.json()] == ["submitted", "failed"]

    listed = await client.get("/v1/runs")
    saved = next(item for item in listed.json() if item["id"] == run_ids[0])
    assert saved["latestOutcome"]["outcome"] == "failed"
    assert saved["latestOutcome"]["reasonCode"] == "validation_failed"

    statistics = await client.get("/v1/application-statistics")
    assert statistics.status_code == 200
    assert statistics.json() == {
        "totalRuns": 2,
        "resolvedRuns": 1,
        "pendingRuns": 1,
        "submitted": 0,
        "failed": 1,
        "cancelled": 0,
        "abandoned": 0,
        "outcomeUncertain": 0,
        "resolutionRate": 0.5,
        "submittedRate": 0.0,
        "generatedAt": statistics.json()["generatedAt"],
    }

    missing_run = await client.put(
        f"/v1/runs/{uuid4()}/outcome",
        json={
            "outcome": "cancelled",
            "reason_code": "user_cancelled",
            "confirmed_by_user": True,
        },
    )
    assert missing_run.status_code == 404


async def test_submission_requires_authorization(client: httpx.AsyncClient) -> None:
    created = await client.post(
        "/v1/runs",
        json={
            "job_id": str(uuid4()),
            "candidate_profile_id": str(uuid4()),
            "candidate_profile_version": 1,
            "auto_submit_authorized": False,
            "job_url": "https://example.com/jobs/2",
            "start_immediately": False,
        },
    )
    run_id = created.json()["id"]
    states = [
        "ingesting_job",
        "preparing_materials",
        "opening_application",
        "filling",
        "validating",
        "ready_to_submit",
    ]
    for index, state in enumerate(states):
        response = await client.post(
            f"/v1/runs/{run_id}/transitions",
            json={
                "to_state": state,
                "reason_code": "test",
                "idempotency_key": f"transition-test-key-{index:04d}",
            },
        )
        assert response.status_code == 200

    blocked = await client.post(
        f"/v1/runs/{run_id}/transitions",
        json={
            "to_state": "submitting",
            "reason_code": "test",
            "idempotency_key": "transition-test-key-submit",
        },
    )
    assert blocked.status_code == 409


async def test_run_controls_checkpoint_pause_resume_and_cancel(
    client: httpx.AsyncClient,
) -> None:
    class BrowserSocket:
        def __init__(self) -> None:
            self.messages: list[dict[str, object]] = []

        async def send_json(self, message: dict[str, object]) -> None:
            self.messages.append(message)

    app = client._transport.app  # type: ignore[attr-defined]
    service = app.state.service
    socket = BrowserSocket()
    service.browser_socket = socket
    service.browser_worker_connected = True

    created = await client.post(
        "/v1/runs",
        json={
            "job_id": str(uuid4()),
            "candidate_profile_id": str(uuid4()),
            "candidate_profile_version": 1,
            "auto_submit_authorized": False,
            "job_url": "https://example.com/jobs/controlled",
            "start_immediately": False,
        },
    )
    run_id = created.json()["id"]
    pause_request = {
        "command": "pause",
        "idempotency_key": "user-pause-control-0001",
    }
    paused = await client.post(f"/v1/runs/{run_id}/control", json=pause_request)
    assert paused.status_code == 200
    assert paused.json()["state"] == "paused"
    assert socket.messages[-1]["action"] == "pause"

    checkpoint = await client.get(f"/v1/runs/{run_id}/checkpoint")
    assert checkpoint.status_code == 200
    assert checkpoint.json()["state"] == "created"
    assert checkpoint.json()["sequence"] == 1

    repeated_pause = await client.post(f"/v1/runs/{run_id}/control", json=pause_request)
    assert repeated_pause.status_code == 200
    assert repeated_pause.json()["state"] == "paused"
    assert len(socket.messages) == 1

    stale_automation = await client.post(
        f"/v1/runs/{run_id}/transitions",
        json={
            "to_state": "filling",
            "reason_code": "late_browser_result",
            "idempotency_key": "late-browser-result-0001",
        },
    )
    assert stale_automation.status_code == 409

    resumed = await client.post(
        f"/v1/runs/{run_id}/control",
        json={
            "command": "resume",
            "idempotency_key": "user-resume-control-0001",
        },
    )
    assert resumed.status_code == 200
    assert resumed.json()["state"] == "created"
    assert socket.messages[-1]["action"] == "resume"

    cancelled = await client.post(
        f"/v1/runs/{run_id}/control",
        json={
            "command": "cancel",
            "idempotency_key": "user-cancel-control-0001",
        },
    )
    assert cancelled.status_code == 200
    assert cancelled.json()["state"] == "cancelled"
    assert cancelled.json()["latestOutcome"]["outcome"] == "cancelled"
    assert socket.messages[-1]["action"] == "cancel"

    cannot_resume = await client.post(
        f"/v1/runs/{run_id}/control",
        json={
            "command": "resume",
            "idempotency_key": "user-resume-control-0002",
        },
    )
    assert cannot_resume.status_code == 409

    events = await client.get(f"/v1/runs/{run_id}/events")
    assert [event["toState"] for event in events.json()] == ["paused", "created", "cancelled"]


async def test_synthetic_demo_creates_durable_run_and_starts_typed_scan(
    client: httpx.AsyncClient, tmp_path: Path
) -> None:
    class BrowserSocket:
        def __init__(self) -> None:
            self.messages: list[dict[str, object]] = []

        async def send_json(self, message: dict[str, object]) -> None:
            self.messages.append(message)

    created = await client.put(
        "/v1/profile",
        json={
            "expected_version": None,
            "profile": {
                "first_name": "PrivateFirst",
                "last_name": "Candidate",
                "email": "synthetic@example.com",
                "github_url": "https://github.com/synthetic-candidate",
                "work_authorization": "authorized",
            },
        },
    )
    profile = created.json()["profile"]
    app = client._transport.app  # type: ignore[attr-defined]
    service = app.state.service
    socket = BrowserSocket()
    service.browser_socket = socket
    service.browser_worker_connected = True

    started = await client.post(
        "/v1/demo/synthetic-run",
        json={
            "candidate_profile_id": profile["id"],
            "candidate_profile_version": profile["version"],
        },
    )

    assert started.status_code == 201
    run = started.json()
    assert run["platform"] == "synthetic"
    assert run["state"] == "opening_application"
    assert run["autoSubmitAuthorized"] is False
    assert socket.messages[0]["action"] == "open_synthetic_form"
    assert socket.messages[0]["envelope"]["authorizationScope"] == "inspect"  # type: ignore[index]

    await service.handle_observed_form(
        run_id=UUID(run["id"]),
        action_id=socket.messages[0]["commandId"],
        observed_form=ObservedForm.model_validate(
            {
                "page_url": "https://synthetic.careerflow.invalid/application",
                "page_state_hash": "abc12345",
                "controls": [
                    {
                        "control_id": "first-name",
                        "label": "First name",
                        "kind": "text",
                        "required": True,
                        "sensitivity": "ordinary",
                        "autocomplete": "given-name",
                    },
                    {
                        "control_id": "email",
                        "label": "Email address",
                        "kind": "email",
                        "required": True,
                        "sensitivity": "sensitive",
                        "autocomplete": "email",
                    },
                ],
            }
        ),
    )
    fill_command = socket.messages[1]
    assert fill_command["action"] == "fill"
    assert fill_command["envelope"]["authorizationScope"] == "fill"  # type: ignore[index]
    assert [item["canonicalPath"] for item in fill_command["fills"]] == [  # type: ignore[index]
        "identity.first_name"
    ]
    assert fill_command["blockedMappings"][0]["canonicalPath"] == "contact.email"  # type: ignore[index]

    await service.handle_fill_result(
        run_id=UUID(run["id"]),
        action_id=fill_command["commandId"],
        ok=True,
        page_state_hash="def67890",
        filled_mappings=[
            FieldMapping(
                control_id="first-name",
                canonical_path="identity.first_name",
                source="deterministic",
                confidence=0.99,
                sensitivity="ordinary",
                decision="auto_fill",
                rationale="Autocomplete metadata identifies identity.first_name.",
            )
        ],
        blocked_mappings=[
            FieldMapping.model_validate(fill_command["blockedMappings"][0])  # type: ignore[index]
        ],
    )
    history = await client.get(f"/v1/runs/{run['id']}/events")
    assert history.status_code == 200
    assert history.json()[-1]["toState"] == "awaiting_human"
    assert history.json()[-2]["safeDetails"] == {
        "filled_control_count": 1,
        "review_required_count": 1,
    }
    explanations = await client.get(f"/v1/runs/{run['id']}/field-explanations")
    assert explanations.status_code == 200
    decisions = {item["canonicalPath"]: item for item in explanations.json()}
    assert decisions["identity.first_name"]["filled"] is True
    assert decisions["contact.email"]["filled"] is False
    assert decisions["contact.email"]["decision"] == "human_required"
    assert "PrivateFirst" not in explanations.text
    assert "synthetic@example.com" not in explanations.text
    original_ids = [item["id"] for item in explanations.json()]

    await service.handle_fill_result(
        run_id=UUID(run["id"]),
        action_id=fill_command["commandId"],
        ok=True,
        page_state_hash="def67890",
        filled_mappings=[
            FieldMapping(
                control_id="first-name",
                canonical_path="identity.first_name",
                source="deterministic",
                confidence=0.99,
                sensitivity="ordinary",
                decision="auto_fill",
                rationale="Autocomplete metadata identifies identity.first_name.",
            )
        ],
        blocked_mappings=[
            FieldMapping.model_validate(fill_command["blockedMappings"][0])  # type: ignore[index]
        ],
    )
    repeated = await client.get(f"/v1/runs/{run['id']}/field-explanations")
    assert [item["id"] for item in repeated.json()] == original_ids
    database_bytes = (tmp_path / "careerflow.db").read_bytes()
    assert b"synthetic@example.com" not in database_bytes
    assert b"PrivateFirst" not in database_bytes
